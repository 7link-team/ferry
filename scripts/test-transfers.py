#!/usr/bin/env python3
"""CSV/XLSX real HTTP acceptance. Requires openpyxl; supports all three databases."""
import http.cookiejar
import json
import os
from pathlib import Path
import socket
import subprocess
import tempfile
import time
import urllib.error
import urllib.request

checks = 0
def check(condition, label):
    global checks
    assert condition, label
    checks += 1
    print('PASS', label, flush=True)

def client():
    return urllib.request.build_opener(urllib.request.ProxyHandler({}), urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))

with tempfile.TemporaryDirectory(prefix='platform-test-') as directory:
    data = Path(directory)
    with socket.socket() as sock:
        sock.bind(('127.0.0.1', 0)); port = sock.getsockname()[1]
    base = f'http://127.0.0.1:{port}/api'
    env = dict(os.environ, FERRY_DATA_DIR=directory, FERRY_BIND=f'127.0.0.1:{port}')
    env.pop('FERRY_DATABASE_URL', None)
    if os.environ.get('FERRY_TEST_DATABASE_URL'): env['FERRY_DATABASE_URL'] = os.environ['FERRY_TEST_DATABASE_URL']
    binary = os.environ.get('FERRY_BINARY', 'target/debug/ferry-server')
    def request(c, path, body=None, method=None, expected=200):
        req = urllib.request.Request(base + path, data=json.dumps(body).encode() if body is not None else None,
            method=method, headers={'Content-Type': 'application/json', 'X-Ferry-Request': '1'})
        try: response = c.open(req, timeout=30)
        except urllib.error.HTTPError as e: response = e
        raw = response.read()
        try: value = json.loads(raw) if raw else None
        except ValueError: raise AssertionError((path, response.code, raw))
        if expected is not None: assert response.code == expected, (path, response.code, value)
        return (response.code, value) if expected is None else value
    anonymous = client()
    log = open(data / 'server.log', 'wb')
    def launch():
        proc = subprocess.Popen([binary], env=env, stdout=log, stderr=log)
        for _ in range(150):
            if proc.poll() is not None: raise AssertionError((data / 'server.log').read_text())
            try:
                if request(anonymous, '/status')['ok']: return proc
            except OSError: pass
            time.sleep(.1)
        raise AssertionError('server startup timed out')
    proc = launch()
    try:
        admin, owner, guest, outsider = client(), client(), client(), client()
        password = 'temporary-workspace-password'
        request(admin, '/setup', {'email':'admin@example.invalid','name':'Admin','password':password,'token':(data/'setup-token').read_text().strip()})
        for c, name in [(owner,'Owner'),(guest,'Guest'),(outsider,'Outsider')]:
            request(c, '/users', {'name':name,'email':name.lower()+'@example.invalid','password':password})
        workspace_tenant = request(owner, '/workspaces', {'name':'Tenant','slug':'tenant'})['id']
        workspace_other = request(outsider, '/workspaces', {'name':'Other','slug':'other'})['id']
        root = f'/workspaces/{workspace_tenant}'
        invitation=request(owner,root+'/members/invitation')['code']
        request(guest,'/workspace-memberships',{'code':invitation})
        import csv, io, sqlite3, zipfile
        from concurrent.futures import ThreadPoolExecutor
        import openpyxl
        def upload(resource, rows=None, raw=None, ext='csv', c=owner, status=200):
            resource = {'hosts':'nodes', 'host-groups':'node-groups'}.get(resource, resource)
            if raw is None:
                out=io.StringIO(); writer=csv.DictWriter(out,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows);raw=out.getvalue().encode()
            boundary='transfer-test-boundary'
            body=(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="data.{ext}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode()+raw+f'\r\n--{boundary}--\r\n'.encode())
            req=urllib.request.Request(base+root+'/'+resource+'/imports',data=body,headers={'Content-Type':'multipart/form-data; boundary='+boundary,'X-Ferry-Request':'1'})
            try: response=c.open(req,timeout=90)
            except urllib.error.HTTPError as e: response=e
            data=json.loads(response.read());assert response.code==status,(response.code,data)
            return data
        def download(resource,kind='export',fmt='csv',query='',c=owner,status=200):
            resource = {'hosts':'nodes', 'host-groups':'node-groups'}.get(resource, resource)
            kind = {'export':'exports', 'template':'import-template'}.get(kind, kind)
            req=urllib.request.Request(base+root+'/'+resource+'/'+kind+'?format='+fmt+query,headers={'X-Ferry-Request':'1'})
            try:response=c.open(req,timeout=90)
            except urllib.error.HTTPError as e:response=e
            data=response.read();assert response.code==status,(response.code,data)
            return data
        def commit(resource,preview,c=owner,status=200):
            resource = {'hosts':'nodes', 'host-groups':'node-groups'}.get(resource, resource)
            return request(c,root+'/'+resource+'/imports/'+preview['id']+'/commits',{},expected=status)
        def xlsx(rows):
            book=openpyxl.Workbook();sheet=book.active
            for row in rows:sheet.append(row)
            out=io.BytesIO();book.save(out);return out.getvalue()
        def count(resource):return request(owner,root+'/'+{'hosts':'nodes', 'host-groups':'node-groups'}.get(resource,resource))['total']
        first=upload('hosts',[{'name':'CSV host','connection':'runner','description':'first\nsecond,"quoted"'}])
        check(first['valid'] and first['rows'][0]['row']==2 and count('hosts')==0,'CSV preview validates without committing hosts')
        check(count('activity')==0,'preview rolls back creation audit')
        result=commit('hosts',first);host=result['items'][0]['id']
        check(count('hosts')==1 and commit('hosts',first)==result,'commit creates once and repeat is idempotent')
        check(request(owner,root+'/nodes/'+host)['description']=='first\nsecond,"quoted"','CSV Unicode/quotes/newlines round trip')
        invalid=upload('hosts',[{'name':'valid but rolled back','connection':'runner'},{'name':'bad','connection':'missing'}])
        check(not invalid['valid'] and invalid['rows'][1]['error'] and invalid['id'] is None and count('hosts')==1,'invalid row prevents pending batch and all business mutations')
        typed=upload('hosts',[{'name':'Invalid typed','connection':'ssh','address':'localhost','username':'root','port':'not-a-number'}])
        check(not typed['valid'] and typed['rows'][0]['row']==2 and typed['rows'][0]['error'],'typed errors are attached to original row')
        bad=upload('hosts',raw=b'name,name\na,b\n',status=400)
        check('表头' in bad['error'],'duplicate headers rejected')
        upload('hosts',raw=b'name,unknown\na,b\n',status=400)
        check(True,'unknown headers rejected')
        parsed=upload('tags',raw=xlsx([['name','color'],['XLSX 标签','#123456']]),ext='xlsx')
        check(parsed['valid'],'genuine XLSX parses Unicode values')
        commit('tags',parsed)
        upload('tags',raw=xlsx([['name'],['=1+1']]),ext='xlsx',status=400)
        check(True,'uploaded XLSX formulas rejected')
        book=openpyxl.Workbook();book.active['A1']='name';book.active['A1048576']='distant';out=io.BytesIO();book.save(out)
        upload('tags',raw=out.getvalue(),ext='xlsx',status=400)
        check(True,'sparse XLSX oversized range rejected before allocation')
        upload('tags',raw=b'name\n'+b'row\n'*1001,status=400)
        check(True,'1000-row bound enforced')
        secret='transfer-secret-never-in-preview'
        pending=upload('credentials',[{'name':'API secret','kind':'token','secret':secret}])
        check(pending['valid'] and secret not in json.dumps(pending) and count('credentials')==0,'credential preview masks value and rolls back')
        if not os.environ.get('FERRY_TEST_DATABASE_URL'):
            database=data/'ferry.sqlite';conn=sqlite3.connect(database)
            stored=conn.execute('SELECT payload FROM resource_imports WHERE id=?',(pending['id'],)).fetchone()[0]
            check(secret not in stored and len(stored)>30,'pending import stored encrypted')
        credential=commit('credentials',pending)['items'][0]['id']
        for fmt in ['csv','xlsx']:
            raw=download('credentials',fmt=fmt)
            if fmt=='xlsx':
                book=openpyxl.load_workbook(io.BytesIO(raw));text=str(list(book.active.values))
            else:text=raw.decode('utf-8-sig')
            check(secret not in text and 'API secret' in text,fmt+' export never reveals credential secrets')
        cancel=upload('tags',[{'name':'Cancelled'}]);request(owner,root+'/tags/imports/'+cancel['id'],method='DELETE',expected=204)
        commit('tags',cancel,status=404)
        check(count('tags')==1,'cancel removes pending data and commit cannot replay it')
        dup=upload('variables',[{'name':'duplicate','kind':'secret','value':'a'},{'name':'duplicate','kind':'secret','value':'b'}])
        check(not dup['valid'] and dup['rows'][1]['error'] and count('variables')==0,'preview catches in-batch duplicates through shared validation')
        atomic=upload('variables',[{'name':'batch-first','kind':'plain','value':'one'},{'name':'batch-conflict','kind':'plain','value':'two'}])
        request(owner,root+'/variables',{'name':'batch-conflict','kind':'plain','value':'already'},expected=201)
        commit('variables',atomic,status=409)
        check(request(owner,root+'/variables?q=batch-first')['total']==0,'commit revalidates and rolls back preceding rows on conflict')
        permission=upload('tags',[{'name':'Actor bound'}]);commit('tags',permission,c=guest,status=403)
        request(outsider,f'/workspaces/{workspace_other}/tags/imports/'+permission['id']+'/commits',{},expected=404)
        upload('tags',[{'name':'Denied'}],c=guest,status=403)
        check(secret.encode() not in download('credentials',c=guest),'read-only member export uses the same masked projection as list')
        check(True,'import ownership, workspace isolation and export permissions enforced')
        # An authorized second actor in the same workspace must not consume the first actor's upload.
        role=request(owner,root+'/roles',{'name':'Importer','permissions':['tags.read','tags.create']},expected=200)['id']
        members=request(owner,root+'/members')['items'];member=next(m for m in members if m['name']=='Guest')
        request(owner,root+'/members/'+member['id'],{'version':member['version'],'role_id':role,'state':'active'},method='PATCH')
        download('credentials',c=guest,status=403)
        commit('tags',permission,c=guest,status=404)
        own=upload('tags',[{'name':'Revoked'}],c=guest)
        request(owner,root+'/roles/'+role,{'name':'Importer','permissions':['tags.read'],'version':1},method='PATCH')
        commit('tags',own,c=guest,status=403)
        check(True,'same-workspace actor binding and permission revocation checked at commit')
        rows=[{'name':f'Paged {i:03}','color':''} for i in range(205)]
        commit('tags',upload('tags',rows))
        csv_rows=list(csv.DictReader(io.StringIO(download('tags',query='&q=Paged&sort=-name').decode('utf-8-sig'))))
        check(len(csv_rows)==205 and csv_rows[0]['name']=='Paged 204' and csv_rows[-1]['name']=='Paged 000','export follows filters/sort across multiple pages')
        request(owner,root+'/tags',{'name':'=danger','color':''},expected=201)
        raw=download('tags',query='&q=danger').decode('utf-8-sig')
        check("'=danger" in raw,'CSV export escapes formula-leading text')
        workbook=openpyxl.load_workbook(io.BytesIO(download('tags',fmt='xlsx',query='&q=danger')))
        check(all(cell.data_type!='f' for row in workbook.active for cell in row),'XLSX export writes literal string cells')
        # All ordinary creation resources share the same import path.
        manifest={'apiVersion':'ferry/v1','kind':'Application','metadata':{'name':'web'},'services':[{'name':'web','image':'busybox:1.36'}]}
        definition={'workdir':'/tmp/import-test','steps':[{'name':'check','run':'true'}]}
        fixtures={
            'host-groups':{'name':'Imported group'},
            'pipeline-groups':{'name':'Imported pipeline group'},
            'pipeline-environments':{'name':'Imported environment','target_kind':'host','target_ref':host},
            'repositories':{'name':'External','url':'https://github.com/example/repository.git','credential_id':credential},
            'pipelines':{'name':'Imported pipeline','target_kind':'host','target_ref':host,'definition':json.dumps(definition)},
            'deployments':{'name':'Imported application','mode':'compose','target_kind':'host','target_ref':host,'workdir':'/tmp/import-app','manifest':json.dumps(manifest)},
            'cron-jobs':{'name':'Disabled cron','scope':'host','schedule':'0 * * * *','timezone':'UTC','enabled':'false','action':'command','target_kind':'host','target_ref':host,'command':'true','timeout_seconds':'30'},
            'alert-rules':{'name':'Offline rule','target_kind':'host','target_ref':host,'metric':'offline','enabled':'false'},
            'dictionaries':{'key':'imported','name':'Imported dictionary','items':'[{"value":"a","label":"A"}]'},
            'roles':{'name':'Imported role','permissions':'["tags.read"]'},
            'app-templates':{'name':'Imported app template','content':json.dumps(manifest),'parameters':'[]'},
            'pipeline-templates':{'name':'Imported pipeline template','content':json.dumps(definition),'parameters':'[]'},
            'config-templates':{'name':'Imported file template','content':'hello','parameters':'[]'},
        }
        for resource,body in fixtures.items():
            previous=count(resource);p=upload(resource,[body]);assert p['valid'],(resource,p)
            assert count(resource)==previous
            commit(resource,p);check(count(resource)==previous+1,resource+' import preserves normal creation contract')
        source_url='https://example.invalid/team/source.git'
        connection=request(owner,root+'/code-connections',{'name':'Transfer source','provider':'git','url':source_url},expected=201)['id']
        actions_yaml='# retained comment\njobs:\n  build:\n    runs-on: self-hosted\n    steps:\n      - id: source\n        uses: actions/checkout@v4\n      - run: echo imported\n'
        actions_body={'name':'Actions export','definition_format':'github-actions','definition':actions_yaml,'target_kind':'host','target_ref':host,'executor_bindings':[{'job_id':'build','target_kind':'host','target_ref':host,'image':'host'}],'graph_layout':{'build':{'x':12,'y':34}},'checkout_bindings':[{'job_id':'build','step_id':'source','connection_id':connection,'repository_url':source_url,'ref':'release/non-default','path':'source'}]}
        classification_group=request(owner,root+'/pipeline-groups')['items'][0]['id']
        classification_environment=request(owner,root+'/pipeline-environments')['items'][0]['id']
        classification_app=request(owner,root+'/deployments')['items'][0]['id']
        classification_tag=request(owner,root+'/tags',{'name':'Pipeline classification'},expected=201)['id']
        actions_body.update(group_id=classification_group,application_id=classification_app,tag_ids=[classification_tag],environment_ids=[classification_environment])
        request(owner,root+'/pipelines',actions_body,expected=201)
        exported=list(csv.DictReader(io.StringIO(download('pipelines',query='&q=Actions%20export').decode('utf-8-sig'))))[0]
        check(exported.get('definition_format')=='github-actions','Actions CSV export retains definition format')
        check(json.loads(exported['graph_layout'])==actions_body['graph_layout'],'Actions export preserves canvas layout')
        check(json.loads(exported['checkout_bindings'])==actions_body['checkout_bindings'],'Actions CSV export retains connection, source ref and checkout path')
        exported={key:value for key,value in exported.items() if key not in ('id','version','created_at','updated_at','group_name','application_name','tags','environments')}
        exported['name']='Actions imported'
        imported=upload('pipelines',[exported]);check(imported['valid'],'Actions export can be imported without converting YAML to legacy JSON');commit('pipelines',imported)
        copied=request(owner,root+'/pipelines?q=Actions%20imported')['items'][0]
        check(copied['definition']==actions_yaml and copied['executor_bindings']==actions_body['executor_bindings'],'Actions import preserves original YAML and executor bindings')
        check(copied['checkout_bindings']==actions_body['checkout_bindings'],'Actions CSV import roundtrip preserves complete checkout binding')
        workbook=openpyxl.load_workbook(io.BytesIO(download('pipelines',fmt='xlsx',query='&q=Actions%20export')))
        sheet_rows=list(workbook.active.values);headers=list(sheet_rows[0]);values=dict(zip(headers,sheet_rows[1]))
        values={key:value for key,value in values.items() if key not in ('id','version','created_at','updated_at','group_name','application_name','tags','environments')}
        values['name']='Actions XLSX imported'
        imported_xlsx=upload('pipelines',raw=xlsx([list(values),list(values.values())]),ext='xlsx')
        check(imported_xlsx['valid'],'Actions XLSX with checkout binding can be imported');commit('pipelines',imported_xlsx)
        copied_xlsx=request(owner,root+'/pipelines?q=Actions%20XLSX%20imported')['items'][0]
        check(copied_xlsx['checkout_bindings']==actions_body['checkout_bindings'],'Actions XLSX roundtrip preserves complete checkout binding')
        check(all(copied[key]==actions_body[key] for key in ['group_id','application_id','tag_ids','environment_ids']),'Actions CSV roundtrip preserves group, application, tags and environments')
        secondary=request(owner,root+'/nodes',{'name':'Actions-only target','connection':'runner'},expected=201)['id']
        scoped=dict(actions_body,name='Secondary executor',executor_bindings=[{'job_id':'build','target_kind':'host','target_ref':secondary,'image':'host'}])
        request(owner,root+'/pipelines',scoped,expected=201)
        request(owner,root+'/nodes/'+secondary,method='DELETE',expected=409)
        check(True,'host referenced only by an Actions job cannot be deleted')
        repo=request(owner,root+'/repositories')['items'][0]['id']
        p=upload('marketplace-sources',[{'name':'Imported source','repository_id':repo,'provider':'github','api_base':'https://api.github.com','project':'example/repository','branch':'main','index_path':'index.json'}])
        assert p['valid'],p;commit('marketplace-sources',p)
        check(count('marketplace-sources')==1,'marketplace import stores definition without network sync')
        for resource in [*fixtures,'hosts','tags','credentials','variables','marketplace-sources']:
            book=openpyxl.load_workbook(io.BytesIO(download(resource,'template','xlsx')))
            check(book.sheetnames==['data','字段说明'] and book.active.max_row==1,resource+' downloadable template includes typed field guide')
        for resource in ['activity','members','settings','marketplace','alert-events','cron-runs','pipeline-runs','pipeline-artifacts','deployment-revisions']:
            download(resource);check(True,resource+' read-only export works')
        if not os.environ.get('FERRY_TEST_DATABASE_URL'):
            stored=conn.execute('SELECT payload FROM resource_imports WHERE id=?',(pending['id'],)).fetchone()[0]
            check(stored=='','consumed import clears encrypted payload')
            conn.close()
        # Concurrent confirmation still creates only once.
        concurrent=upload('tags',[{'name':'Concurrent'}])
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(lambda _:commit('tags',concurrent),range(2)))
        check(results[0]==results[1] and request(owner,root+'/tags?q=Concurrent')['total']==1,'concurrent commit is serialized and idempotent')
        request(owner,root+'/variables',{'name':'a-secret','kind':'secret','value':'hidden-export-value'},expected=201)
        request(owner,root+'/variables',{'name':'z-plain','kind':'plain','value':'visible-export-value'},expected=201)
        exported=download('variables').decode('utf-8-sig')
        check('visible-export-value' in exported and 'hidden-export-value' not in exported,'secret first row never omits later plain variable values from export')
        tag=request(owner,root+'/tags',{'name':'Production'},expected=201)['id']
        tagged=upload('hosts',[{'name':'Tagged host','connection':'runner','tag_ids':json.dumps([tag])}])
        assert tagged['valid'],tagged
        hid=commit('hosts',tagged)['items'][0]['id']
        saved=request(owner,root+'/nodes/'+hid)
        check(saved['tag_ids']==[tag] and saved['tags'][0]['name']=='Production','host tags import persists association and readable labels')
        filtered=request(owner,root+'/nodes?tag_id='+tag)
        check(filtered['total']==1 and filtered['items'][0]['id']==hid,'host tag filtering shares pagination and tenant scope')
        check('Tagged host' in download('hosts',query='&tag_id='+tag).decode('utf-8-sig'),'host export preserves tag filter')
        request(owner,root+'/tags/'+tag,method='DELETE',expected=409)
        old=dict(saved);old.pop('tag_ids');request(owner,root+'/nodes/'+hid,old,method='PATCH')
        check(request(owner,root+'/nodes/'+hid)['tag_ids']==[tag],'legacy host updates preserve omitted tag associations')
        request(owner,root+'/nodes/'+hid,dict(saved,tag_ids=[]),method='PATCH',expected=409)
        check(request(owner,root+'/nodes/'+hid)['tag_ids']==[tag],'stale host edit cannot clear tag associations')
        foreign=request(outsider,f'/workspaces/{workspace_other}/tags',{'name':'Foreign'},expected=201)['id']
        current=request(owner,root+'/nodes/'+hid)
        request(owner,root+'/nodes/'+hid,dict(current,tag_ids=[foreign]),method='PATCH',expected=400)
        check(request(owner,root+'/nodes/'+hid)==current,'foreign tag rejection rolls back the whole host update')
        request(owner,root+'/nodes/'+hid,dict(current,tag_ids=[tag,tag]),method='PATCH',expected=400)
        request(owner,root+'/nodes/'+hid,dict(current,tag_ids=[]),method='PATCH')
        check(request(owner,root+'/nodes/'+hid)['tag_ids']==[],'explicit empty tag array clears association')
        request(owner,root+'/tags/'+tag,method='DELETE',expected=204)
        tag=request(owner,root+'/tags',{'name':'Cleanup'},expected=201)['id']
        current=request(owner,root+'/nodes/'+hid)
        request(owner,root+'/nodes/'+hid,dict(current,tag_ids=[tag]),method='PATCH')
        request(owner,root+'/nodes',{'ids':[hid]},method='DELETE',expected=200)
        request(owner,root+'/tags/'+tag,method='DELETE',expected=204)
        check(True,'host bulk deletion removes associations so tags can be deleted')
        parent=request(owner,root+'/node-groups',{'name':'Parent'},expected=201)['id']
        child=request(owner,root+'/node-groups',{'name':'Child','parent_id':parent},expected=201)['id']
        foreign=request(outsider,f'/workspaces/{workspace_other}/node-groups',{'name':'Other parent'},expected=201)['id']
        request(owner,root+'/node-groups',{'name':'Wrong parent','parent_id':foreign},expected=400)
        request(owner,root+'/node-groups/'+parent,{'name':'Parent','parent_id':child,'version':1},method='PATCH',expected=400)
        request(owner,root+'/node-groups/'+parent,method='DELETE',expected=409)
        request(owner,root+'/node-groups/'+child,method='DELETE',expected=204)
        request(owner,root+'/node-groups/'+parent,method='DELETE',expected=204)
        check(True,'group ancestry rejects foreign parents/cycles and protects children')
        large=upload('config-templates',[{'name':'Large first','content':'a'*180000,'parameters':'[]'},{'name':'Large second','content':'b'*180000,'parameters':'[]'}])
        check(large['valid'] and large['count']==2,'multipart upload above normal JSON body limit is supported within file limit')
        request(owner,root+'/config-templates/imports/'+large['id'],method='DELETE',expected=204)
        print(f'{checks} checks passed',flush=True)
    finally:
        proc.terminate(); proc.wait(timeout=20); log.close()
